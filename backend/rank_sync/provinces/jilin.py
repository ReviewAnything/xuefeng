"""吉林：jleea.com.cn「高考成绩一分段表」（十年一行的宽表；xls → 2025 起 PDF）。

版式（五年同一张表，只是载体换了）：

- **每行一个整十分档**（`700`、`690` …），右侧 10 列是 `+9 … +0`，
  格子里是**累计人数**（不是本段人数）——所以本段人数要靠「本行 − 右邻」反算；
- 表体左侧留白 = 该分无人：顶部分档因此天然是截断桶（官方通稿写作
  「704-750 | 11 | 11」），首个有值的格子就是「X 分及以上」；
- 表体一直印到 **0 分**（0 分那一格是缺考/零分考生，如 2025 物理 1323 人），
  所以 `score_floor=0`、`total` = 0 分累计；
- 2022/2023 是文理分科（理工/文史），2024 起 3+1+2（物理-普通/历史-普通），
  落盘统一映射为 物理类/历史类 并在 `meta.subject_aliases` 留痕。

两个载体坑：

1. **2024 起附件名仍写 `.xls`，真身是 xlsx**（PK 头）。`openpyxl` 见到 `.xls`
   后缀会直接拒（`InvalidFileException`），必须喂 `BytesIO`；2022/2023 才是
   真 BIFF（`xlrd`）。所以按魔数分派，绝不看后缀。
2. **2023 正文的 `<a>` 嵌套写坏了**（`<a>` 里又开 `<a>`，且少闭合）：
   理工那一行的链接被劈成「…含照顾分」和「)」两个锚点，后面的语文/体育锚点
   还嵌在里头。用正则抓 `<a ...>(.*?)</a>` 会把体育的表名粘到文史行上，
   所以改用栈式解析——**遇到新 `<a>` 先闭合上一个**，与浏览器行为一致。

附件名规则：同一篇通告挂 6 个附件，只有前两个是普通类（理工|物理 / 文史|历史），
其余是 体育 / 艺术（同名另发），按标签里的「体育」「艺术」挡掉。
"""
from __future__ import annotations

import io
import json
import re
from html.parser import HTMLParser

from ..common import http_get, now_iso, sha256, validate_rows

PROVINCE = "吉林"
SITE = "https://www.jleea.com.cn"
SCHEME = "3+1+2"
SUBJECTS = ["物理类", "历史类"]
STATUS = "ready"

# 站内两个前端接口：详情给正文（含附件锚点），列表给栏目分页
DETAIL_API = SITE + "/server-front/front/content/detail"
LIST_API = SITE + "/server-front/front/content/page"
NOTICE_CHANNEL = "ksdt"      # 通知公告（与 SOURCES 里登记的公告页同栏目）
LIST_PAGE_SIZE = 100
LIST_MAX_PAGES = 30

# 每年：公告页（可点开源溯）+ 两份普通类附件
SOURCES = {
    2022: {"notice": SITE + "/front/content/4172",
           "files": {"物理类": SITE + "/u/cms/www/uploadfile/2022/0623/20220623023325476.xls",
                     "历史类": SITE + "/u/cms/www/uploadfile/2022/0623/20220623024117265.xls"}},
    2023: {"notice": SITE + "/front/content/4546",
           "files": {"物理类": SITE + "/u/cms/www/uploadfile/2023/0623/20230623115245271.xls",
                     "历史类": SITE + "/u/cms/www/uploadfile/2023/0623/20230623115422738.xls"}},
    2024: {"notice": SITE + "/front/content/4864",
           "files": {"物理类": SITE + "/u/cms/www/uploadfile/2024/0625/20240625112434246.xls",
                     "历史类": SITE + "/u/cms/www/uploadfile/2024/0625/20240625112449739.xls"}},
    2025: {"notice": SITE + "/front/content/201719",
           "files": {"物理类": SITE + "/u/cms/www/2025/06/25/1937708593876869123.pdf",
                     "历史类": SITE + "/u/cms/www/2025/06/25/1937708686835228675.pdf"}},
    2026: {"notice": SITE + "/front/content/202730",
           "files": {"物理类": SITE + "/u/cms/www/2026/06/25/2069968279602028546.pdf",
                     "历史类": SITE + "/u/cms/www/2026/06/25/2069968368907149314.pdf"}},
}

# 2022/2023 官方仍是文理分科，落盘映射 + 留痕
ALIASES = {2022: {"历史类": "文史", "物理类": "理工"},
           2023: {"历史类": "文史", "物理类": "理工"}}

# 表名守卫：只认普通类（理工/文史 或 物理-普通/历史-普通），体育/艺术表另发
TITLE_PAT = re.compile(
    r"^(20\d{2})年吉林省普通高校招生考试(理工|文史|物理-普通|历史-普通)1分段表")
# 附件/标题里的「另发表」关键词（体育、艺术、高职对口）
AUX_PAT = re.compile(r"体育|艺术|对口|高职|专升本")
HEADER_PAT = re.compile(r"^\+([0-9])$")
SCORE_LABEL_PAT = re.compile(r"^\d{1,3}$")
BAND_WIDTH = 10      # 每行 10 个 1 分档（+9 … +0）


# ---------------------------------------------------------------------------
# 正文锚点：栈式解析（2023 的 a 嵌套写坏了）
# ---------------------------------------------------------------------------
class _AnchorStack(HTMLParser):
    """把正文里所有 `<a>` 收成 [(href, label)]。

    遇到新的 `<a>` 就先闭合上一个未闭合的——浏览器就这么干，2023 那篇通告
    依赖这个行为才能把「理工 / 文史」的表名和体育、艺术的表名分开。
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[tuple[str, str]] = []
        self._href: str | None = None
        self._buf: list[str] = []

    def _flush(self) -> None:
        if self._href is not None:
            self.out.append((self._href, re.sub(r"\s+", " ", "".join(self._buf)).strip()))
        self._href, self._buf = None, []

    def handle_starttag(self, tag, attrs):  # noqa: D102 - HTMLParser 钩子
        if tag == "a":
            self._flush()
            self._href = dict(attrs).get("href") or ""
        elif tag == "br":
            self._buf.append(" ")

    def handle_endtag(self, tag):  # noqa: D102
        if tag == "a":
            self._flush()

    def handle_data(self, data):  # noqa: D102
        if self._href is not None:
            self._buf.append(data)


def parse_anchors(body: str) -> list[tuple[str, str]]:
    """正文 HTML -> [(href, label)]（纯函数）。"""
    p = _AnchorStack()
    p.feed(body or "")
    p.close()
    p._flush()
    return p.out


# 科类 -> 附件标签关键词（理工/文史 是 2022–2023 的写法，物理/历史 是 2024 起）
LABEL_KEYS = {"物理类": ("理工", "物理"), "历史类": ("文史", "历史")}


def _abs_url(href: str) -> str:
    if href.startswith("//"):
        return "https:" + href
    if href.startswith("/"):
        return SITE + href
    if href.startswith("http"):
        return href
    return SITE + "/" + href


def pick_files(pairs: list[tuple[str, str]]) -> dict[str, str]:
    """从锚点里挑出两份普通类附件（纯函数）。

    2023 那篇通告会多出一个标签为「)」的锚点（理工链接被劈开的残片，指向体育(理)
    的文件），它不含科类关键词，自然被跳过。
    """
    out: dict[str, str] = {}
    for subj in SUBJECTS:
        for href, label in pairs:
            if not any(k in label for k in LABEL_KEYS[subj]):
                continue
            if AUX_PAT.search(label):
                continue
            out[subj] = _abs_url(href)
            break
    missing = [s for s in SUBJECTS if s not in out]
    if missing:
        raise ValueError(f"吉林公告页缺附件：{missing}")
    return out


# ---------------------------------------------------------------------------
# 宽表 -> 行
# ---------------------------------------------------------------------------
def _num(v) -> int | None:
    """单元格 -> 整数。兼容 float、'692'、带千分位。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip().replace(",", "")
    return int(s) if s.isdigit() else None


def bands_to_rows(bands: list[tuple[int, dict[int, int]]]) -> list[list[int]]:
    """宽表 -> [[score, 本段, 累计]]（纯函数，唯一的口径转换点）。

    ``bands`` 是 `(整十分档, {偏移: 累计人数})`，偏移 9 = 该档 +9 分。同一档内
    分数高的在前，档与档之间按分数降序接续。

    **首个有值的格子按「X 分及以上」处理**（`本段 = 累计`）：官方表在最高分那档
    只印到有人为止，左侧留白既可能是「该分无人」也可能是「与右侧同值省略不印」，
    两种读法给出同一个累计数，所以取「首个格子即截断桶」——不臆造空白格的分数。
    """
    if not bands:
        raise ValueError("未解析出任何分档行")
    flat: list[tuple[int, int]] = []
    for base, cells in bands:
        if base % BAND_WIDTH:
            raise ValueError(f"分档不是整十：{base}")
        if not cells:
            continue
        if len(cells) > BAND_WIDTH:
            raise ValueError(f"分档 {base} 有 {len(cells)} 列，超过 {BAND_WIDTH}")
        for off in sorted(cells, reverse=True):
            flat.append((base + off, cells[off]))
    rows: list[list[int]] = []
    prev_cum = prev_score = None
    for score, cum in flat:
        if prev_score is not None and score >= prev_score:
            raise ValueError(f"分档重叠或非降序：{prev_score} -> {score}")
        seg = cum if prev_cum is None else cum - prev_cum
        if seg < 0:
            raise ValueError(f"累计人数回落 score={score}：{prev_cum} -> {cum}")
        rows.append([score, seg, cum])
        prev_cum, prev_score = cum, score
    validate_rows(rows)
    return rows


def _table_of_xls(blob: bytes) -> list[list]:
    """xls/xlsx 统一读成二维表：按魔数分派，不看后缀。

    2024 起的附件叫 `.xls` 其实是 xlsx（PK 头）；`openpyxl.load_workbook` 只
    认文件名时会因后缀拒收，所以一律走 `BytesIO`。
    """
    if blob[:2] == b"PK":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if blob[:4] == b"\xd0\xcf\x11\xe0":
        import xlrd

        sh = xlrd.open_workbook(file_contents=blob).sheet_by_index(0)
        return [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    raise ValueError(f"既不是 BIFF 也不是 OOXML（前 4 字节 {blob[:4]!r}）")


def parse_table(table: list[list]) -> tuple[list[list[int]], str]:
    """二维表 -> (rows, title)（纯函数，xls/xlsx 共用）。"""
    hdr = None
    for i, row in enumerate(table[:6]):
        cells = [str(v).strip() for v in row if v is not None]
        if cells and "分数" in cells[0] and \
                sum(bool(HEADER_PAT.match(c)) for c in cells) >= BAND_WIDTH - 2:
            hdr = i
            break
    if hdr is None:
        raise ValueError("吉林 xls 未找到「分数 +9…+0」表头，版式可能变了")
    title = str(table[0][0]).strip() if table and table[0] else ""
    cmap: dict[int, int] = {}     # 列号 -> 偏移
    for c, v in enumerate(table[hdr]):
        m = HEADER_PAT.match(str(v).strip())
        if m:
            cmap[c] = int(m.group(1))
    if len(cmap) != BAND_WIDTH:
        raise ValueError(f"表头列数 {len(cmap)} != {BAND_WIDTH}")
    bands = []
    for row in table[hdr + 1:]:
        base = _num(row[0]) if row else None
        if base is None:
            continue
        cells = {}
        for c, off in cmap.items():
            v = _num(row[c]) if c < len(row) else None
            if v is not None:
                cells[off] = v
        if cells:
            bands.append((base, cells))
    return bands_to_rows(bands), title


def _page_groups(page, tol: float = 3.0) -> list[list[tuple[float, str]]]:
    """一页 -> 按 y 拼好的行（每行按 x 升序）。

    同一行的数字 y 基本一致（行距 ~24pt），用「首个成员的 y」做锚，
    免得相邻行的边角字被吸进来。
    """
    from ..common import page_lines

    anchors: list[float] = []
    buckets: list[list[tuple[float, float, str]]] = []
    for x, y, t in sorted(page_lines(page), key=lambda r: -r[1]):
        if anchors and abs(y - anchors[-1]) <= tol:
            buckets[-1].append((x, y, t))
        else:
            anchors.append(y)
            buckets.append([(x, y, t)])
    return [[(x, t) for x, _y, t in sorted(b, key=lambda c: c[0])] for b in buckets]


def _columns_of(nums: list[tuple[float, int]]) -> list[int]:
    """分档行 -> 每个数所属列序（0 = `+9` 列，9 = `+0` 列）。

    两条独立判据必须一致，不一致就是版式漂了，整份拒收：

    1. **右侧满列**：官方表只在最高分那档左侧留白，其余满 10 列，所以 n 个数的
       列序必然是 `10-n … 9`；
    2. **按 x 反推**：同一行的数字右对齐、列距恒定，用中位列距把 x 换回列序。

    判据 2 能抓住的是**中间留白**（某档跳了一列，如 5/6/8/9 而不是 6/7/8/9）——
    那会让分数整体错位一分。左侧留白不是异常：它正是「X 分及以上」截断桶，
    两种判据在此时同解。

    只印一个数的行直接拒收——那个数属于哪一分无从判断，不猜。
    """
    n = len(nums)
    if n < 2:
        raise ValueError("分档行只印了 1 个数，无法判定它属于哪一分")
    count_idx = list(range(BAND_WIDTH - n, BAND_WIDTH))
    xs = [x for x, _ in nums]
    diffs = sorted(b - a for a, b in zip(xs, xs[1:]))
    spacing = diffs[len(diffs) // 2]
    if spacing <= 0:
        raise ValueError(f"列距异常：{spacing}")
    last = xs[-1]
    x_idx = [BAND_WIDTH - 1 - round((last - x) / spacing) for x in xs]
    if x_idx != count_idx:
        raise ValueError(f"列归属不一致：按 x={x_idx}，按列数={count_idx}")
    return count_idx


def parse_pdf(blob: bytes) -> tuple[list[list[int]], str]:
    """PDF（与 xls 同一张宽表）-> (rows, title)。"""
    import tempfile
    from pathlib import Path

    from pdfminer.high_level import extract_pages

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        f.write(blob)
        path = f.name
    bands: list[tuple[int, dict[int, int]]] = []
    title = ""
    try:
        for page in extract_pages(path):
            for cells in _page_groups(page):
                texts = [t for _, t in cells]
                joined = " ".join(texts)
                if "分段表" in joined and not SCORE_LABEL_PAT.match(texts[0]):
                    title = title or joined
                    continue
                if "分数" in texts[0]:        # 每页重复的表头
                    continue
                if not SCORE_LABEL_PAT.match(texts[0]):
                    continue
                nums = [(x, int(t)) for x, t in cells[1:] if t.isdigit()]
                if not nums:
                    continue
                base = int(texts[0])
                offsets = {BAND_WIDTH - 1 - idx: v for idx, (_, v) in
                           zip(_columns_of(nums), nums)}
                bands.append((base, offsets))
    finally:
        Path(path).unlink(missing_ok=True)
    return bands_to_rows(bands), title


def parse_attachment(blob: bytes) -> tuple[list[list[int]], str]:
    """按魔数分派：%PDF / PK(xlsx) / OLE2(xls)。"""
    if blob[:5] == b"%PDF-":
        return parse_pdf(blob)
    return parse_table(_table_of_xls(blob))


# ---------------------------------------------------------------------------
# 发现 / 取数
# ---------------------------------------------------------------------------
def _subject_of_title(title: str) -> str:
    """理工/文史（2022–2023）与 物理-普通/历史-普通（2024 起）都归到两科类。"""
    if "文史" in title or "历史" in title:
        return "历史类"
    if "理工" in title or "物理" in title:
        return "物理类"
    raise ValueError(f"吉林表名无法判定科类：{title!r}")


def _detail_body(content_id: str) -> str:
    """正文 HTML（前端详情接口）。"""
    blob = http_get(f"{DETAIL_API}?id={content_id}", timeout=60)
    data = (json.loads(blob.decode("utf-8")) or {}).get("data") or {}
    items = {it.get("field"): it.get("value")
             for it in (data.get("instance") or {}).get("instanceItems") or []}
    return (items.get("contents") or {}).get("正文") or ""


def _files_from_notice(notice: str) -> dict[str, str]:
    cid = notice.rstrip("/").rsplit("/", 1)[-1]
    if not cid.isdigit():
        raise ValueError(f"吉林公告页 URL 里没有内容 id：{notice}")
    body = _detail_body(cid)
    if not body:
        raise ValueError(f"吉林公告页正文为空：{notice}")
    return pick_files(parse_anchors(body))


def discover(year: int) -> dict:
    """翻「通知公告」列表找该年主表通告（按发布时间倒序，翻到更早的年份就停）。

    注意同一篇通告常同时投到「通知公告」和「考试院公告」，内容一致；
    这里只走通知公告，与 SOURCES 里登记的 URL 保持同一栏目。
    """
    for page_no in range(1, LIST_MAX_PAGES + 1):
        url = (f"{LIST_API}?channelPaths={NOTICE_CHANNEL}"
               f"&pageSize={LIST_PAGE_SIZE}&pageNum={page_no}")
        data = (json.loads(http_get(url, timeout=60).decode("utf-8")) or {}).get("data") or {}
        recs = data.get("records") or []
        if not recs:
            break
        for rec in recs:
            title = (rec.get("title") or "").strip()
            # 主表标题以年份开头、以「高考成绩一分段表」结尾；体育/艺术另有通告
            if not title.startswith(str(year)) or "高考成绩一分段表" not in title:
                continue
            if AUX_PAT.search(title):
                continue
            return {"notice": f"{SITE}/front/content/{rec['id']}"}
        oldest = str(recs[-1].get("publishTime") or "")[:4]
        if oldest.isdigit() and int(oldest) < year:
            break
    raise ValueError(f"吉林 {year} 高考成绩一分段表通告未找到")


def source(year: int) -> dict:
    return SOURCES.get(year) or discover(year)


def parse(year: int) -> dict:
    src = source(year)
    files = dict(src.get("files") or _files_from_notice(src["notice"]))
    subjects: dict = {}
    shas: dict[str, str] = {}
    urls: dict[str, str] = {}
    floors: dict[str, int] = {}
    for subj in SUBJECTS:
        url = files.get(subj)
        if not url:
            raise ValueError(f"吉林 {year} 缺 {subj} 的附件")
        blob = http_get(url, timeout=180)
        rows, title = parse_attachment(blob)
        if not TITLE_PAT.match(title):
            raise ValueError(f"吉林 {year} 附件不是普通类分段表（表名：{title!r}），拒收")
        if _subject_of_title(title) != subj:
            raise ValueError(f"吉林 {year} 科类不符：期望 {subj}，表名 {title!r}")
        subjects[subj] = {
            "rows": rows,
            "top_note": f"{rows[0][0]}↑",     # 首个有值的格子就是「X 分及以上」
            "total": rows[-1][2],
            "table_title": title,
            # 分档行数（十年一行），用于人工核对版式有没有少半页
            "bands": rows[0][0] // BAND_WIDTH - rows[-1][0] // BAND_WIDTH + 1,
        }
        shas[subj] = sha256(blob)
        urls[subj] = url
        floors[subj] = rows[-1][0]
    meta = {
        "year": year, "province": PROVINCE, "data_source": "official",
        "source_url": src["notice"],
        "subject_scheme": "文理分科" if year <= 2023 else SCHEME,
        "columns": 3,
        "band_width": BAND_WIDTH,
        "tracks": {"primary": "高考总分（含照顾分）"},
        "note": ("表体是「整十分档 × 每行 10 个 1 分档」的宽表，格子印的是累计人数，"
                 "本段人数由「本行 − 右邻」反算；顶部左侧留白即「X 分及以上」截断桶；"
                 "表体印到 0 分（含缺考计 0 的考生），故 score_floor=0；"
                 "正文另挂体育/艺术同款表，是不同考生群体，不并入"),
        "attachment_url_map": urls,
        "attachment_sha256_map": shas,
        "attachment_url": urls.get("物理类", ""),
        "attachment_sha256": shas.get("物理类", ""),
        "score_floor": floors,
        "fetched_at": now_iso(),
    }
    if year in ALIASES:
        meta["subject_aliases"] = ALIASES[year]
    return {"meta": meta, "subjects": subjects}
